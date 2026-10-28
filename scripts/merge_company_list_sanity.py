#!/usr/bin/env python3
"""Merge acquired company-list snapshots into the job-digest workbook."""
from __future__ import annotations

import csv
import hashlib
import json
from difflib import SequenceMatcher
import os
import re
import shutil
import tempfile
import unicodedata
import zipfile
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "artifacts/company-list-sanity"
WORKBOOK = ROOT / "artifacts/job-digest-mining-report.xlsx"
SOURCE_SHEET = "5 - Jev Preliminary"
CAREERS_SHEET = "6 - Company Careers"
ALL_SHEET = "7 - US Company Lists"
CROSS_SHEET = "8 - Company Cross-Ref"


def safe_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def key(value):
    text = unicodedata.normalize("NFKD", safe_text(value) or "").encode("ascii", "ignore").decode().casefold()
    text = re.sub(r"&", " and ", text)
    text = re.sub(r"\b(incorporated|inc|corporation|corp|company|co|limited|ltd|llc|plc|lp|llp)\b", " ", text)
    return re.sub(r"[^a-z0-9]", "", text)


def load_lists():
    rows = []
    sec = json.loads((DATA / "sec-company-tickers-exchange.json").read_text())
    for cik, name, ticker, exchange in sec["data"]:
        rows.append(["SEC EDGAR ticker/exchange registrants", "Current at download (2026-10-24)", name, cik, ticker, exchange,
                     "SEC filer/ticker list; includes foreign issuers and is not a U.S.-only registry.", "https://www.sec.gov/files/company_tickers_exchange.json"])
    with (DATA / "sp500-constituents.csv").open(newline="", encoding="utf-8-sig") as f:
        for item in csv.DictReader(f):
            rows.append(["S&P 500 snapshot (DataHub mirror)", "Snapshot downloaded 2026-10-24; embedded source vintage unclear", item.get("Name"), None, item.get("Symbol"), item.get("Sector"),
                         "Third-party snapshot; issuer domicile is not guaranteed U.S.; share classes may duplicate issuers.", "https://datahub.io/core/s-and-p-500-companies-financials/r/constituents.csv"])
    with (DATA / "fortune500-2025-extracted.csv").open(newline="", encoding="utf-8-sig") as f:
        for item in csv.DictReader(f):
            rows.append(["Fortune 2025 ranking page (partial extraction)", "2025", item.get("company_name"), item.get("rank"), None, None,
                         "Unofficial extraction from public page; 996 of 1,000 ranks extracted (4 absent); not complete dataset.", "https://fortune.com/ranking/fortune500/2025/"])
    return rows


def check_xlsx_parts():
    with zipfile.ZipFile(WORKBOOK) as z:
        names = z.namelist()
    risky = [n for n in names if any(part in n.lower() for part in ("drawing", "chart", "media/", "pivot", "slicer", "embedding", "vbaproject", "activex", "customxml"))]
    if risky:
        raise RuntimeError(f"Unsupported non-cell workbook content found; refusing merge: {risky}")


def main():
    if not WORKBOOK.exists():
        raise FileNotFoundError(WORKBOOK)
    check_xlsx_parts()
    original_hash = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    wb = load_workbook(WORKBOOK)
    if SOURCE_SHEET not in wb.sheetnames or CAREERS_SHEET not in wb.sheetnames:
        raise RuntimeError(f"Expected source sheets missing: {wb.sheetnames}")
    if ALL_SHEET in wb.sheetnames or CROSS_SHEET in wb.sheetnames:
        raise RuntimeError("Output sheets already exist; refusing duplicate/overwrite")

    # Source company IDs/names preserve duplicates by ID. The company-careers tab
    # was the enriched workbook's one-row-per-source-ID roster.
    roster = []
    sh = wb[CAREERS_SHEET]
    headers = {safe_text(c.value): i for i, c in enumerate(sh[1])}
    for row in sh.iter_rows(min_row=2, values_only=True):
        roster.append({"id": row[headers["Report company ID"]], "name": row[headers["Observed company name (preserved)"]],
                       "canonical": row[headers["Canonical company name"]], "status": row[headers["Identity status"]]})
    if len(roster) != 13494:
        raise RuntimeError(f"Expected 13,494 source IDs, got {len(roster)}")

    lists = load_lists()
    bykey = defaultdict(list)
    for company in roster:
        normalized = key(company["name"])
        if normalized:
            bykey[normalized].append(company)
    name_keys = list(bykey)
    blocked_keys = defaultdict(list)
    for existing_key in name_keys:
        blocked_keys[existing_key[:4]].append(existing_key)

    all_ws = wb.create_sheet(ALL_SHEET)
    all_headers = ["Source list", "List year / snapshot", "Listed company name", "Rank / CIK", "Ticker", "Sector / exchange", "Coverage / limitations", "Source URL"]
    all_ws.append(all_headers)
    for row in lists:
        all_ws.append(row)

    # Deduplicate source entries by normalized name within each source. Retain
    # all linked source names in the combined tab; cross-reference source facts.
    grouped = defaultdict(list)
    for row in lists:
        source, vintage, name, rank, ticker, sector, caveat, url = row
        normalized = key(name)
        if normalized:
            grouped[(source, normalized)].append(row)

    cross = wb.create_sheet(CROSS_SHEET)
    cross_headers = ["Source list", "List year / snapshot", "Listed company name", "Rank / CIK", "Ticker", "Sector / exchange", "Match type", "Matched source company ID(s)", "Observed company name(s)", "Canonical name(s)", "Current research status(es)", "Review note", "Source URL"]
    cross.append(cross_headers)
    seen = set()
    for (source, normalized), candidates in sorted(grouped.items()):
        row = candidates[0]
        _, vintage, listed, rank, ticker, sector, caveat, url = row
        exact = bykey.get(normalized, [])
        if exact:
            match_type = "Exact normalized name"
            matches = exact
            note = "Name match only; does not independently verify U.S. location or identity. Review entity/brand/subsidiary alignment."
        else:
            # High threshold, and only propose a fuzzy result when it is clearly
            # separated from the runner-up. These remain review-only.
            best_score, best_key, second_score = 0.0, None, 0.0
            for existing_key in blocked_keys.get(normalized[:4], ()):
                score = SequenceMatcher(None, normalized, existing_key).ratio()
                if score > best_score:
                    second_score, best_score, best_key = best_score, score, existing_key
                elif score > second_score:
                    second_score = score
            best = (best_score, bykey.get(best_key, []))
            second = second_score
            if best[0] >= 0.92 and best[0] - second >= 0.04:
                match_type = f"Candidate name match ({best[0]:.0%})"
                matches = best[1]
                note = "Fuzzy suggestion only; manual review required."
            else:
                match_type = "No name match"
                matches = []
                note = "No clear match under conservative name normalization. This does not imply the company is absent."
        ids = sorted({str(m["id"]) for m in matches if m["id"]})
        observed = sorted({str(m["name"]) for m in matches if m["name"]})
        canon = sorted({str(m["canonical"]) for m in matches if m["canonical"]})
        statuses = sorted({str(m["status"]) for m in matches if m["status"]})
        cross.append([source, vintage, listed, rank, ticker, sector, match_type, "; ".join(ids), "; ".join(observed), "; ".join(canon), "; ".join(statuses), note, url])
        seen.add((source, normalized))

    # Workbook styling and usability.
    for ws in (all_ws, cross):
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="214E6B")
        for col in range(1, ws.max_column + 1):
            values = [len(str(ws.cell(r, col).value or "")) for r in range(1, min(ws.max_row, 150) + 1)]
            ws.column_dimensions[get_column_letter(col)].width = min(50, max(12, (max(values) if values else 12) + 2))

    if hashlib.sha256(WORKBOOK.read_bytes()).hexdigest() != original_hash:
        raise RuntimeError("Workbook changed during merge; aborting")
    backup = WORKBOOK.with_suffix(".pre-company-list-sanity.bak.xlsx")
    shutil.copy2(WORKBOOK, backup)
    fd, tmp_name = tempfile.mkstemp(prefix=WORKBOOK.stem + ".", suffix=".xlsx", dir=WORKBOOK.parent)
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        wb.save(tmp)
        check_xlsx_parts()
        check = load_workbook(tmp, read_only=True, data_only=False)
        if check.sheetnames[-2:] != [ALL_SHEET, CROSS_SHEET]:
            raise RuntimeError("Output sheets did not validate")
        if check[ALL_SHEET].max_row != len(lists) + 1:
            raise RuntimeError("Combined sheet row count mismatch")
        if check[CROSS_SHEET].max_row != len(grouped) + 1:
            raise RuntimeError("Cross-reference row count mismatch")
        check.close()
        os.replace(tmp, WORKBOOK)
    finally:
        if tmp.exists():
            tmp.unlink()
    print(json.dumps({"workbook": str(WORKBOOK), "backup": str(backup), "source_rows": len(lists), "distinct_source_names": len(grouped), "cross_reference_rows": len(grouped), "source_company_ids": len(roster), "tabs": [ALL_SHEET, CROSS_SHEET]}, indent=2))

if __name__ == "__main__":
    main()
