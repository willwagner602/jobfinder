from openpyxl import load_workbook
from collections import Counter, defaultdict
import unicodedata, re

path = 'artifacts/job-digest-mining-report.xlsx'
wb = load_workbook(path, read_only=True, data_only=True)

def norm(value):
    s = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().casefold()
    s = re.sub(r'\b(incorporated|inc|corporation|corp|company|co|limited|ltd|llc|plc|lp|llp)\b', ' ', s)
    return re.sub(r'[^a-z0-9]', '', s)

identity = wb['1 - Company Identity']
headers = [c.value for c in next(identity.iter_rows(min_row=1, max_row=1))]
ix = {name: i for i, name in enumerate(headers)}
companies = {row[ix['Report company ID']].value: row[ix['Observed company name']].value for row in identity.iter_rows(min_row=2)}

site = wb['2 - Site Evidence']
headers = [c.value for c in next(site.iter_rows(min_row=1, max_row=1))]
ix = {name: i for i, name in enumerate(headers)}
classes = Counter()
site_ids = defaultdict(set)
for row in site.iter_rows(min_row=2):
    cid = row[ix['Report company ID']].value
    kind = row[ix['URL class']].value
    classes[kind] += 1
    site_ids[kind].add(cid)

lists = wb['7 - US Company Lists']
headers = [c.value for c in next(lists.iter_rows(min_row=1, max_row=1))]
ix = {name: i for i, name in enumerate(headers)}
per_source = defaultdict(set)
all_keys = set()
for row in lists.iter_rows(min_row=2):
    source = row[ix['Source list']].value
    k = norm(row[ix['Listed company name']].value)
    if k:
        all_keys.add(k)
        per_source[source].add(k)

company_keys = defaultdict(list)
for cid, name in companies.items():
    k = norm(name)
    if k:
        company_keys[k].append(cid)
matched_ids = {cid for k in all_keys for cid in company_keys.get(k, [])}
matched_names = {k for k in all_keys if k in company_keys}
site_union = set().union(*site_ids.values())
site_match = site_union.intersection(matched_ids)

print('Original company IDs:', len(companies))
print('Companies with any Site Evidence rows:', len(site_union))
print('Site evidence rows by class:', dict(classes))
print('Distinct company IDs by site class:', {k: len(v) for k, v in site_ids.items()})
print('Distinct normalized names in combined new lists:', len(all_keys))
print('Distinct normalized names per source:', {k: len(v) for k, v in per_source.items()})
print('Original company IDs exact-normalized matching a new list:', len(matched_ids))
print('Distinct matched original names:', len(matched_names))
print('Companies with site evidence matching new list:', len(site_match))
print('Companies with site evidence not matching new list:', len(site_union - site_match))
print('Matches by site class:', {k: len(v.intersection(matched_ids)) for k, v in site_ids.items()})
company_site_ids = site_ids.get('Company domain', set()) | site_ids.get('ATS board', set())
print('Companies with company-domain or ATS-board evidence:', len(company_site_ids))
print('Those with exact normalized match to reference lists:', len(company_site_ids.intersection(matched_ids)))
print('Those without exact normalized match:', len(company_site_ids - matched_ids))
wb.close()
