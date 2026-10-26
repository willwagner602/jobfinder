#!/tmp/.wbvenv/bin/python
import csv, difflib, hashlib, os, re, tempfile, unicodedata
from collections import defaultdict, Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

P=Path('artifacts/job-digest-mining-report.xlsx')
OUT=Path('artifacts/company-list-sanity/company-list-cross-reference.xlsx')
A='A - Job Ad Companies'; B='B - Reference Companies'
def norm(v):
 s=unicodedata.normalize('NFKD',str(v or '')).encode('ascii','ignore').decode().casefold().replace('&',' and ')
 s=re.sub(r'\b(incorporated|inc|corporation|corp|company|co|limited|ltd|llc|plc|lp|llp)\b',' ',s)
 return re.sub(r'[^a-z0-9]','',s)
def headers(ws): return {c.value:i for i,c in enumerate(next(ws.iter_rows(min_row=1,max_row=1)))}
def val(v): return '' if v is None else v

def main():
 digest=hashlib.sha256(P.read_bytes()).hexdigest()
 src=load_workbook(P,read_only=True,data_only=True)
 identity=src['1 - Company Identity']; ih=headers(identity)
 sites=src['2 - Site Evidence']; sh=headers(sites)
 research=src['6 - Company Careers']; rh=headers(research)
 references=src['7 - US Company Lists']; bh=headers(references)
 research_by_id={}
 for r in research.iter_rows(min_row=2,values_only=True):
  research_by_id[str(val(r[rh['Report company ID']]))]={'canonical':val(r[rh['Canonical company name']]),'status':val(r[rh['Identity status']]),'careers':val(r[rh['Careers page / board']]),'official':val(r[rh['Official site']])}
 site_rows=defaultdict(list)
 for r in sites.iter_rows(min_row=2,values_only=True):
  site_rows[str(val(r[sh['Report company ID']]))].append((val(r[sh['URL class']]),val(r[sh['URL / board root']])))
 originals=[]; by_name=defaultdict(list)
 for r in identity.iter_rows(min_row=2,values_only=True):
  cid=str(val(r[ih['Report company ID']])); name=val(r[ih['Observed company name']]); nk=norm(name); ev=site_rows.get(cid,[])
  item={'id':cid,'name':name,'variants':val(r[ih['Observed name variants']]),'jobs':val(r[ih['Unique job URLs']]),'titles':val(r[ih['Distinct raw titles']]),'locations':val(r[ih['Top locations']]),'sources':val(r[ih['Sources']]),'classes':'; '.join(sorted({str(e[0]) for e in ev if e[0]})),'roots':'; '.join(sorted({str(e[1]) for e in ev if e[1]})),'research':research_by_id.get(cid,{}),'key':nk}
  originals.append(item)
  if nk: by_name[nk].append(item)
 refs=[]
 for r in references.iter_rows(min_row=2,values_only=True):
  x={'source':val(r[bh['Source list']]),'vintage':val(r[bh['List year / snapshot']]),'name':val(r[bh['Listed company name']]),'rank':val(r[bh['Rank / CIK']]),'ticker':val(r[bh['Ticker']]),'sector':val(r[bh['Sector / exchange']]),'caveat':val(r[bh['Coverage / limitations']]),'url':val(r[bh['Source URL']])}; x['key']=norm(x['name']); refs.append(x)
 blocks=defaultdict(list)
 for k in by_name: blocks[k[:1]].append(k)
 fuzzy_ids=set()
 for x in refs:
  k=x['key']; matches=by_name.get(k,[])
  if matches: x['match']='Exact normalized name'; x['records']=matches
  else:
   best=0; second=0; bestkey=None
   for cand in blocks.get(k[:1],[]):
    if abs(len(k)-len(cand))>max(4,int(len(k)*.25)): continue
    score=difflib.SequenceMatcher(None,k,cand,autojunk=False).ratio()
    if score>best: second,best,bestkey=best,score,cand
    elif score>second: second=score
   if bestkey and best>=.90 and best-second>=.04: x['match']='Possible name match - review'; x['records']=by_name[bestkey]; fuzzy_ids.update(r['id'] for r in x['records'])
   else: x['match']='No clear name match'; x['records']=[]
  x['ids']=sorted({z['id'] for z in x['records']})
 if hashlib.sha256(P.read_bytes()).hexdigest()!=digest: raise RuntimeError('Source workbook changed during read')
 out=Workbook(); out.remove(out.active)
 aw=out.create_sheet(A); aw.append(['Original company ID','Observed company name','Observed name variants','Unique job URLs','Distinct raw titles','Top locations','Observed sources','Site evidence classes','Site evidence roots','Careers research status','Canonical research name','Careers page / board','Official site','Reference match','Matched reference sources','Reference names','Interpretation'])
 exact_id_sources=defaultdict(set); fuzzy_id_sources=defaultdict(set); id_ref_names=defaultdict(set)
 ref_keys={x['key'] for x in refs if x['key']}
 for x in refs:
  for z in x['records']:
   id_ref_names[z['id']].add(str(x['name']))
   (exact_id_sources if x['match']=='Exact normalized name' else fuzzy_id_sources)[z['id']].add(str(x['source']))
 for x in originals:
  d=x['research']; mt='Exact normalized name' if x['key'] in ref_keys else ('Possible name match - review' if x['id'] in fuzzy_ids else 'No clear name match')
  ss=exact_id_sources[x['id']] if mt=='Exact normalized name' else fuzzy_id_sources[x['id']]
  aw.append([x['id'],x['name'],x['variants'],x['jobs'],x['titles'],x['locations'],x['sources'],x['classes'],x['roots'],d.get('status',''),d.get('canonical',''),d.get('careers',''),d.get('official',''),mt,'; '.join(sorted(ss)),'; '.join(sorted(id_ref_names[x['id']])), 'Name match is not identity or U.S. status verification.' if mt!='No clear name match' else 'No match in acquired lists does not imply absence or invalidity.'])
 bw=out.create_sheet(B); bw.append(['Reference source','List vintage / download','Reference company name','Rank / CIK','Ticker','Sector / exchange','Match result','Original company ID(s)','Original observed name(s)','Canonical research name(s)','Careers research status(es)','Site evidence classes','Coverage / data caveat','Source URL'])
 for x in refs:
  rs=x['records']; ids=x['ids']; states=sorted({str(z['research'].get('status')) for z in rs if z['research'].get('status')}); names=sorted({str(z['name']) for z in rs if z['name']}); can=sorted({str(z['research'].get('canonical')) for z in rs if z['research'].get('canonical')}); cls=sorted({str(e[0]) for z in rs for e in site_rows.get(z['id'],[]) if e[0]})
  bw.append([x['source'],x['vintage'],x['name'],x['rank'],x['ticker'],x['sector'],x['match'],'; '.join(ids),'; '.join(names),'; '.join(can),'; '.join(states),'; '.join(cls),x['caveat'],x['url']])
 for ws in (aw,bw):
  ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
  for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='214E6B')
  for col in range(1,ws.max_column+1):
   lens=[len(str(ws.cell(r,col).value or '')) for r in range(1,min(ws.max_row,150)+1)]
   ws.column_dimensions[get_column_letter(col)].width=min(48,max(12,max(lens,default=12)+2))
 OUT.parent.mkdir(parents=True,exist_ok=True)
 fd,tmp=tempfile.mkstemp(prefix='company-list-cross-reference.',suffix='.xlsx',dir=OUT.parent); os.close(fd)
 try:
  out.save(tmp)
  check=load_workbook(tmp,read_only=True,data_only=True)
  if check[A].max_row!=len(originals)+1 or check[B].max_row!=len(refs)+1: raise RuntimeError('row count validation failed')
  check.close(); os.replace(tmp,OUT)
 finally:
  if os.path.exists(tmp): os.unlink(tmp)
 print('Output',OUT,'A rows',len(originals),'B rows',len(refs),'A matches',dict(Counter(aw.cell(r,14).value for r in range(2,aw.max_row+1))),'B matches',dict(Counter(x['match'] for x in refs)))
 src.close()
if __name__=='__main__': main()
